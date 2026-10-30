"""SYNTHETIC TEST FIXTURES - NOT NISR DATA.

Builds small Excel workbooks that imitate the *layout* of NISR SAS table
releases (title rows, 'District' header, 30 districts, a national row labelled
'National' / 'Total' / 'S/Total', renumbered sheets) so that pipeline behaviour
can be tested deterministically. Every value is invented and every workbook
carries the marker "SYNTHETIC TEST FIXTURE". These files are written only to
pytest temporary directories and are never used by the app.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from src.config import DISTRICTS

MARKER = "SYNTHETIC TEST FIXTURE - not NISR data"
CROPS = ["Maize", "Sorghum", "Paddy rice", "Wheat", "Cassava", "Sweet potato", "Irish potato",
         "Cooking banana", "Bean", "Soybean"]


def _crop_sheet(wb, title, label, values, national_label="National"):
    ws = wb.create_sheet(title)
    ws.append([None, label])
    ws.append([None, "District", "Cereals", *CROPS])
    for d, row in values.items():
        ws.append([None, d if d != "National" else national_label, None, *row])
    ws.append([None, MARKER])


def _input_sheet(wb, title, label, group, values, national_label="National"):
    ws = wb.create_sheet(title)
    ws.append([None, label])
    ws.append([None, "District", group, None, None])
    ws.append([None, None, "Overall", "SSF", "LSF"])
    for d, v in values.items():
        ws.append([None, d if d != "National" else national_label, v, v, None])
    ws.append([None, MARKER])


def make_workbook(path: Path, season_year: int, season: str, yield_scale: float = 1.0,
                  start_table: int = 10, drop_input: bool = False, bad_district: bool = False,
                  low_districts: tuple[str, ...] = ("Ruhango",), seed_override: dict | None = None) -> Path:
    tag = f"{season_year} Season {season}"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    names = list(DISTRICTS)
    if bad_district:
        names[0] = "Atlantis"
    yields, areas = {}, {}
    for i, d in enumerate(names):
        base = 2000.0 if d not in low_districts else 1200.0
        yields[d] = [base * yield_scale + i, 0, 0, 0, 12000, 8000, 0, 14000, 700, 500]
        areas[d] = [1000 + 10 * i, 0, 0, 0, 900, 900, 0, 900, 2000, 600]
    yields["National"] = [2000.0 * yield_scale, 0, 0, 0, 12000, 8000, 9000, 14000, 700, 500]
    areas["National"] = [sum(a[k] for a in areas.values()) for k in range(len(CROPS))]
    areas["National"][6] = 100.0      # national Irish potato row must be observed
    prod = {d: [y * a / 1000 for y, a in zip(yields[d], areas[d])] for d in yields}
    n = start_table
    _crop_sheet(wb, f"Table {n}", f"Table {n}: {tag}_Cultivated area by crop type and district (Ha)", areas); n += 1
    _crop_sheet(wb, f"Table {n}", f"Table {n}: {tag}_Harvested area by crop type and district (Ha)", areas); n += 1
    _crop_sheet(wb, f"Table {n}", f"Table {n}: Season {season} {season_year}_Average yield by crop type and district (Kg/Ha)", yields); n += 1
    _crop_sheet(wb, f"Table {n}", f"Table {n}: Season {season} {season_year}_Crop production by crop type and district (MT)",
                prod, national_label="S/Total"); n += 1
    seed = {d: (20.0 if d in low_districts else 45.0) for d in names}
    seed["National"] = 37.0
    if seed_override:
        seed.update(seed_override)
    fert = {d: (30.0 if d in low_districts else 70.0) for d in names}
    fert["National"] = 60.0
    _input_sheet(wb, f"Table {n}", f"Table {n}: {tag}_Use of seeds by farmer type per district (Percentage)",
                 "Percentage of farmers who used improved seeds", seed); n += 1
    if not drop_input:
        _input_sheet(wb, f"Table {n}", f"Table {n}: {tag}_Use of Inorganic fertilizer by farmer type per district (Percentage)",
                     "Percentage of farmers who applied inorganic fertilizer", fert, national_label="Total"); n += 1
    _input_sheet(wb, f"Table {n}", f"Table {n}: {tag}_Use of organic fertilizer by farmer type per district (Percentage)",
                 "Percentage of farmers who applied organic fertilizer", fert); n += 1
    path = Path(path)
    wb.save(path)
    return path
